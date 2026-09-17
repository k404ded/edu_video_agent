"""
Generates editable PowerPoint presentations (.pptx) from structured educational slide content.

Features:
- 16:9 widescreen layout
- Professional visual styling with title slide & section cards
- Formatted on-screen bullet points and key term tags
- Embedded speaker notes containing the exact voice-over script for every slide
"""
import os
from pathlib import Path
from typing import Optional, List
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from models.schema import PipelineResult, SlideOutput, DocumentSummary


# Color Palette
COLOR_BG = RGBColor(248, 250, 252)         # Slate 50
COLOR_CARD_BG = RGBColor(255, 255, 255)    # Pure White
COLOR_BORDER = RGBColor(226, 232, 240)     # Slate 200
COLOR_PRIMARY = RGBColor(15, 23, 42)       # Slate 900
COLOR_SECONDARY = RGBColor(71, 85, 105)    # Slate 600
COLOR_ACCENT = RGBColor(37, 99, 235)       # Blue 600
COLOR_TAG_BG = RGBColor(239, 246, 255)     # Blue 50
COLOR_TAG_TEXT = RGBColor(29, 78, 216)     # Blue 700
COLOR_ACTION_BG = RGBColor(254, 243, 199)  # Amber 100
COLOR_ACTION_TEXT = RGBColor(180, 83, 9)   # Amber 700


def _add_card(slide, left, top, width, height, fill_color=COLOR_CARD_BG, border_color=COLOR_BORDER):
    """Adds a clean card container with optional border."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def _set_slide_background(slide, width, height, color=COLOR_BG):
    """Fills slide background with a smooth solid color."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, width, height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg


def generate_presentation(result: PipelineResult, output_dir: str = "output/generated") -> str:
    """
    Creates a 16:9 editable PowerPoint presentation matching the PipelineResult.
    Embeds slide voiceover into speaker notes for each slide.
    Returns the absolute path to the generated presentation.pptx.
    """
    os.makedirs(output_dir, exist_ok=True)
    out_path = os.path.abspath(os.path.join(output_dir, "presentation.pptx"))

    prs = Presentation()
    # 16:9 Widescreen (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # -------------------------------------------------------------------------
    # 1. Title Slide
    # -------------------------------------------------------------------------
    title_slide = prs.slides.add_slide(blank_layout)
    _set_slide_background(title_slide, prs.slide_width, prs.slide_height, COLOR_BG)

    # Decorative header bar
    top_bar = title_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_ACCENT
    top_bar.line.fill.background()

    # Title Card Container
    _add_card(title_slide, Inches(1.2), Inches(1.2), Inches(10.933), Inches(5.1))

    # Badge: Topic / Course
    badge_box = title_slide.shapes.add_textbox(Inches(1.6), Inches(1.6), Inches(10.1), Inches(0.4))
    tf_badge = badge_box.text_frame
    tf_badge.word_wrap = True
    p_badge = tf_badge.paragraphs[0]
    p_badge.text = "EDUCATIONAL PRESENTATION"
    p_badge.font.size = Pt(11)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_ACCENT

    # Main Title
    title_box = title_slide.shapes.add_textbox(Inches(1.6), Inches(2.1), Inches(10.1), Inches(1.6))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = result.document_summary.main_topic or "Educational Overview"
    p_title.font.size = Pt(36)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_PRIMARY

    # Learning Objective
    obj_box = title_slide.shapes.add_textbox(Inches(1.6), Inches(3.8), Inches(10.1), Inches(1.3))
    tf_obj = obj_box.text_frame
    tf_obj.word_wrap = True
    p_obj_lbl = tf_obj.paragraphs[0]
    p_obj_lbl.text = "🎯 Learning Objective:"
    p_obj_lbl.font.size = Pt(13)
    p_obj_lbl.font.bold = True
    p_obj_lbl.font.color.rgb = COLOR_SECONDARY

    p_obj = tf_obj.add_paragraph()
    p_obj.text = result.document_summary.learning_objective or "Understand core concepts and principles."
    p_obj.font.size = Pt(15)
    p_obj.font.color.rgb = COLOR_PRIMARY

    # Metadata footer: Target Audience
    if result.document_summary.audience_level:
        aud_box = title_slide.shapes.add_textbox(Inches(1.6), Inches(5.3), Inches(10.1), Inches(0.5))
        tf_aud = aud_box.text_frame
        p_aud = tf_aud.paragraphs[0]
        p_aud.text = f"Audience Level: {result.document_summary.audience_level}"
        p_aud.font.size = Pt(11)
        p_aud.font.color.rgb = COLOR_SECONDARY

    # Speaker notes for title slide
    title_notes = title_slide.notes_slide.notes_text_frame
    title_notes.text = (
        f"Welcome. Today we will explore {result.document_summary.main_topic}. "
        f"By the end of this session: {result.document_summary.learning_objective}"
    )

    # -------------------------------------------------------------------------
    # 2. Content Slides (One per SlideOutput)
    # -------------------------------------------------------------------------
    total_slides = len(result.slide_outputs)
    for idx, slide_out in enumerate(result.slide_outputs, start=1):
        slide = prs.slides.add_slide(blank_layout)
        _set_slide_background(slide, prs.slide_width, prs.slide_height, COLOR_BG)

        # Header bar
        header_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.1))
        header_bar.fill.solid()
        header_bar.fill.fore_color.rgb = COLOR_ACCENT
        header_bar.line.fill.background()

        # Section counter & Title
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
        tf_h = header_box.text_frame
        tf_h.word_wrap = True
        p_sub = tf_h.paragraphs[0]
        p_sub.text = f"SECTION {idx} OF {total_slides}"
        p_sub.font.size = Pt(10)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_ACCENT

        p_h = tf_h.add_paragraph()
        p_h.text = slide_out.section_title or f"Section {idx}"
        p_h.font.size = Pt(24)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_PRIMARY

        # Left Card: Main On-Screen Content
        _add_card(slide, Inches(0.8), Inches(1.5), Inches(7.5), Inches(5.2))
        content_box = slide.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(6.9), Inches(4.7))
        tf_content = content_box.text_frame
        tf_content.word_wrap = True

        p_c_lbl = tf_content.paragraphs[0]
        p_c_lbl.text = "Key Concepts & Information"
        p_c_lbl.font.size = Pt(13)
        p_c_lbl.font.bold = True
        p_c_lbl.font.color.rgb = COLOR_SECONDARY

        # Parse bullets or text lines
        raw_lines = [line.strip("- *• \t") for line in slide_out.on_screen_content.split("\n") if line.strip()]
        if not raw_lines:
            raw_lines = [slide_out.on_screen_content]

        for line in raw_lines:
            p_line = tf_content.add_paragraph()
            p_line.text = f"• {line}"
            p_line.font.size = Pt(14)
            p_line.font.color.rgb = COLOR_PRIMARY
            p_line.space_after = Pt(8)

        # Right Column Top: Learning Objective Card
        _add_card(slide, Inches(8.6), Inches(1.5), Inches(3.9), Inches(2.2))
        lo_box = slide.shapes.add_textbox(Inches(8.8), Inches(1.7), Inches(3.5), Inches(1.8))
        tf_lo = lo_box.text_frame
        tf_lo.word_wrap = True
        p_lo_lbl = tf_lo.paragraphs[0]
        p_lo_lbl.text = "🎯 Objective"
        p_lo_lbl.font.size = Pt(12)
        p_lo_lbl.font.bold = True
        p_lo_lbl.font.color.rgb = COLOR_ACCENT

        p_lo_val = tf_lo.add_paragraph()
        p_lo_val.text = slide_out.learning_objective or "Understand core section concepts."
        p_lo_val.font.size = Pt(12)
        p_lo_val.font.color.rgb = COLOR_PRIMARY
        p_lo_val.space_before = Pt(4)

        # Right Column Bottom: Key Terms & Visual Cue Card
        _add_card(slide, Inches(8.6), Inches(3.9), Inches(3.9), Inches(2.8))
        aux_box = slide.shapes.add_textbox(Inches(8.8), Inches(4.0), Inches(3.5), Inches(2.5))
        tf_aux = aux_box.text_frame
        tf_aux.word_wrap = True

        # Key terms
        p_kt_lbl = tf_aux.paragraphs[0]
        p_kt_lbl.text = "🔑 Key Terms"
        p_kt_lbl.font.size = Pt(11)
        p_kt_lbl.font.bold = True
        p_kt_lbl.font.color.rgb = COLOR_SECONDARY

        terms_str = ", ".join(slide_out.key_terms) if slide_out.key_terms else "(general concepts)"
        p_kt_val = tf_aux.add_paragraph()
        p_kt_val.text = terms_str
        p_kt_val.font.size = Pt(11)
        p_kt_val.font.color.rgb = COLOR_TAG_TEXT
        p_kt_val.space_after = Pt(8)

        # Visual Action Guide
        p_va_lbl = tf_aux.add_paragraph()
        p_va_lbl.text = "🎬 Visual Action"
        p_va_lbl.font.size = Pt(11)
        p_va_lbl.font.bold = True
        p_va_lbl.font.color.rgb = COLOR_ACTION_TEXT

        p_va_val = tf_aux.add_paragraph()
        p_va_val.text = slide_out.visual_action or "Present and emphasize key points."
        p_va_val.font.size = Pt(11)
        p_va_val.font.color.rgb = COLOR_PRIMARY

        # ---------------------------------------------------------------------
        # Embedded Speaker Notes (Exact Voice-Over Narration)
        # ---------------------------------------------------------------------
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = slide_out.voice_over

    prs.save(out_path)
    return out_path
