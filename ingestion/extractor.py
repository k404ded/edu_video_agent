"""
Document ingestion layer.

Responsible ONLY for turning an uploaded file (or pasted text) into a
list of normalized Section objects. No LLM calls happen here — this
module is pure extraction, so the "source of truth" text is captured
before any AI reasoning touches it.
"""
from pathlib import Path
from typing import List

from models.schema import Section


def extract_pptx(file_path: str) -> List[Section]:
    from pptx import Presentation

    prs = Presentation(file_path)
    sections = []
    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        text_chunks = []

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            # Heuristic: the title placeholder (or the first text box) becomes the title
            if shape == slide.shapes.title or (not title and shape.shape_id == slide.shapes[0].shape_id):
                if not title:
                    title = text.split("\n")[0]
            text_chunks.append(text)

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        raw_text = "\n".join(text_chunks).strip()
        if not title:
            title = f"Slide {i}"

        sections.append(Section(index=i, title=title, raw_text=raw_text, notes=notes))

    return sections


def extract_pdf(file_path: str) -> List[Section]:
    import pdfplumber

    sections = []
    with pdfplumber.open(file_path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            first_line = text.split("\n")[0] if text else f"Page {i}"
            title = first_line[:80] if first_line else f"Page {i}"
            sections.append(Section(index=i, title=title, raw_text=text))
    return sections


def extract_text(raw_text: str, chunk_by_blank_lines: bool = True) -> List[Section]:
    """
    Split pasted plain text into sections. Splits on double-newlines
    (paragraph/section breaks) so behavior is predictable and the user
    can control chunking simply by how they paste/format the text.
    """
    raw_text = raw_text.strip()
    if not raw_text:
        return []

    if chunk_by_blank_lines:
        chunks = [c.strip() for c in raw_text.split("\n\n") if c.strip()]
    else:
        chunks = [raw_text]

    sections = []
    for i, chunk in enumerate(chunks, start=1):
        first_line = chunk.split("\n")[0]
        title = first_line[:80] if first_line else f"Section {i}"
        sections.append(Section(index=i, title=title, raw_text=chunk))
    return sections


def extract_sections(file_path: str = None, pasted_text: str = None) -> List[Section]:
    """
    Single entry point used by the orchestrator. Exactly one of
    file_path or pasted_text should be provided.
    """
    if pasted_text and pasted_text.strip():
        return extract_text(pasted_text)

    if not file_path:
        raise ValueError("Provide either a file_path or pasted_text.")

    ext = Path(file_path).suffix.lower()
    if ext == ".pptx":
        return extract_pptx(file_path)
    elif ext == ".pdf":
        return extract_pdf(file_path)
    elif ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return extract_text(f.read())
    else:
        raise ValueError(f"Unsupported file type: {ext}")
